#!/usr/bin/env python3
"""Reproduce every table, statistic and figure.   Usage:  python scripts/run_all.py"""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from at_mmc import data as D, derived as X, stats as S, plotting as P, reproduction as R
from at_mmc.config import TABLES, RESULTS


def main():
    TABLES.mkdir(parents=True, exist_ok=True)
    tabs = {
        "T03_mass_balance": D.mass_balance(D.load_mix()),
        "T10_correlations": X.table10(),
        "T10x_full_pearson_matrix": X.correlation_matrix().reset_index(names="property"),
        "T10y_full_spearman_matrix": X.correlation_matrix("spearman").reset_index(names="property"),
        "T11_oxide_consistency": X.oxide_consistency(),
        "T12_binder_indices": X.binder_indices(),
        "T13_pct_change_vs_C0": X.pct_change_vs_c0().reset_index(),
        "T14_fibre_series_changes": X.fibre_series_changes().reset_index(),
        "T15_T21_ryshkewitch_fits": S.ryshkewitch_fits(),
        "T16_density_diffusivity_checks": X.physical_thermal_consistency().reset_index(),
        "T17_lca_relative": X.lca_relative().reset_index(),
        "T18_fibre_dose_fits": S.fibre_dose_fits(),
        "T19_two_way_unreplicated": S.two_way_unreplicated(),
        "T19b_kendall_friedman": pd.DataFrame([S.kendall_w_and_friedman()]),
        "T20_mechanical_change": X.mechanical_change_table().reset_index(),
        "T22_contrasts_2x2": S.contrasts_2x2(),
        "T23_replicate_precision": S.replicate_precision(),
        "S1_pareto_dominance": S.pareto_dominance(),
        "S2_desirability_scenarios": S.desirability_scenarios(),
        "S3_reproduction_check": R.compare(),
    }
    for name, df in tabs.items():
        df.to_csv(TABLES / f"{name}.csv", index=False)
    xl = RESULTS / "AT_MMC_results_tables.xlsx"
    with pd.ExcelWriter(xl, engine="openpyxl") as w:
        readme = pd.DataFrame({"Note": [
            "Export of results/tables/*.csv produced by scripts/run_all.py (values, not formulas).",
            "All numbers derive from mixture-level summaries (one value per mixture). F and p values are NOMINAL / descriptive.",
            "They are not replicate-based inference. See README.md and docs/DATA_ISSUES.md.",
            "Sheet S3_reproduction_check compares each recomputed value with the number printed in the manuscript."]})
        readme.to_excel(w, sheet_name="README", index=False)
        for name, df in tabs.items():
            df.to_excel(w, sheet_name=name[:31], index=False)
        for ws in w.book.worksheets:
            for c in ws[1]:
                c.font = Font(name="Arial", bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F4E79")
                c.alignment = Alignment(wrap_text=True, vertical="center")
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    c.font = Font(name="Arial", size=10)
                    if isinstance(c.value, float): c.number_format = "0.0000" if abs(c.value) < 10 else "0.00"
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = min(42, max(12, max(len(str(c.value or "")) for c in col[:30]) + 2))
            ws.freeze_panes = "A2"
    figs = P.make_all()
    chk = tabs["S3_reproduction_check"]
    print(f"Tables written: {len(tabs)}   Figures written: {len(figs)} (PNG + PDF)")
    print(f"Reproduction check: {(chk.status == 'match').sum()}/{len(chk)} manuscript values reproduced")


if __name__ == "__main__":
    main()
